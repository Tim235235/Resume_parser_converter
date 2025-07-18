import re
import ast
import spacy
from docxtpl import DocxTemplate
from docx import Document
import pdfplumber
import pypandoc
import os

# === PDF extraction ===
def extract_text_from_pdf(uploaded_file):
    text = ""
    if uploaded_file:
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    return text

# === Section extraction from text ===
def extract_sections(text):
    pattern = re.compile(
        r"^(%s)\s*$" % "|".join(re.escape(h) for h in all_headings),
        re.MULTILINE | re.IGNORECASE
    )

    heading_map = {}
    for canonical, synonyms in section_synonyms.items():
        for synonym in synonyms:
            heading_map[synonym.lower()] = canonical

    matches = list(pattern.finditer(text.lower().strip()))
    sections = {}
    for i, match in enumerate(matches):
        matched_heading = match.group(1).strip().lower()
        canonical_heading = heading_map.get(matched_heading)
        if not canonical_heading:
            continue  # skip unknown headings
        start = match.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections[canonical_heading] = text[start:end]
    name(text, sections)
    return sections

# === Formatting preview ===
def formatting(path):
    sections = extract_sections(extract_text_from_pdf(path))
    for heading, content in sections.items():
        if content.strip() == "\n":
            content = ""
        print(f"|{heading}|{content}")

# === Name extractor ===
def name(text, info_dic):
    doc = nlp(text.lower())
    for ent in doc.ents:
        if ent.label_ == "PERSON":
            info_dic["Name"] = ent.text
            return
    info_dic["Name"] = "No Name found"

# === Pattern-based finders ===
def age_finder(text):
    for pattern in loaded_patterns.get("age_patterns", []):
        match = re.search(pattern, text)
        if match:
            for group in match.groups():
                if group and group.isdigit():
                    return group
    return "N/A"

def nationality_finder(text):
    doc = nlp(text)
    for pattern in loaded_patterns.get("nationality_patterns", []):
        match = re.search(pattern, text)
        if match:
            nationality = match.group(1)
            if nationality:
                return nationality
    for sent in doc.sents:
        if any(keyword in sent.text.lower() for keyword in ["citizen", "nationality", "citizenship", "passport", "born", "from"]):
            for ent in sent.ents:
                if ent.label_ == "NORP":
                    return ent.text
    return "N/A"

def years_of_experience_finder(text):
    for pattern in loaded_patterns.get("years_experience_patterns", []):
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return "N/A"

def availability_finder(text):
    for pattern in loaded_patterns.get("availability_patterns", []):
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return "N/A"

# === IT skills extraction ===
def it_section(text):
    it_skills_sorted = sorted(loaded_patterns.get("it_skills", []), key=len, reverse=True)
    pattern = r'|'.join(
        fr'(?<!\w){re.escape(skill)}(?!\w)' for skill in it_skills_sorted
    )
    matches = re.findall(pattern, text, re.IGNORECASE)
    return list({m.lower() for m in matches})

# === Hidden data collection ===
def hidden_data_collector(text, section_info):
    section_info["age"] = age_finder(text)
    section_info["nationality"] = nationality_finder(text)
    section_info["years_of_experience"] = years_of_experience_finder(text)
    section_info["availability"] = availability_finder(text)
    it = it_section(text)
    section_info["it"] = "\n\n".join(it)
    return section_info

# === Word document reading ===
def read_word(filename):
    doc = Document(filename)
    full_text = ""
    for paragraph in doc.paragraphs:
        full_text += paragraph.text + "\n"
    return full_text

# === Data population for template ===
def populate_word_dic(dti, ed):
    for key in keys:
        value = dti.get(key)
        ed[key] = value if value is not None else "N/A"
    return ed

# === Replace placeholders in DOCX ===
def replace_placeholders(use_data):
    doc = DocxTemplate(template_docx)
    doc.render(use_data)
    doc.save(filled_template_path)
    return doc

# === Bullet points fix ===
def bullet_points_check(text):
    bullet_pattern = r"[•\-\*\▪\●\‣\∙]\s*\n\s*"
    for key, value in text.items():
        if isinstance(value, str):
            value = re.sub(bullet_pattern, r'\n• ', value)
            text[key] = value
    return text

# === Convert DOCX to PDF ===
def convert_docx_to_pdf():
    base_dir = os.path.dirname(__file__)
    input_docx = os.path.join(base_dir, "filled_template.docx")
    output_pdf = os.path.join(base_dir, "Converted_resume.pdf")
    
    try:
        pypandoc.convert_file(
            input_docx,
            "pdf",
            outputfile=output_pdf
        )
        if os.path.exists(output_pdf):
            return output_pdf
        else:
            print("Conversion reported success but PDF file not found.")
            return None
    except Exception as e:
        print("PDF conversion failed:", e)
        return None



# === Global setup ===
base_dir = os.path.dirname(os.path.abspath(__file__))

template_docx = os.path.join(base_dir, "Eng_TEMPLATE CV IOTA.docx")
filled_template_path = os.path.join(base_dir, "filled_template.docx")
converted_pdf_path = os.path.join(base_dir, "Converted_resume.pdf")

keys = [
    "Name", "age", "nationality", "years_of_experience", "availability",
    "summary", "education", "training", "it", "languages", "professional_experience"
]

nlp = spacy.load('en_core_web_lg')

# Load pattern data files
with open(os.path.join(base_dir, "Section_synonyms_text_file.txt"), "r", encoding="utf-8") as file:
    content = file.read()
section_synonyms = ast.literal_eval(content)

loaded_patterns = {}
with open(os.path.join(base_dir, "Regex_patterns.txt"), "r", encoding="utf-8") as file:
    exec(file.read(), {}, loaded_patterns)

all_headings = [h for synonyms in section_synonyms.values() for h in synonyms]

